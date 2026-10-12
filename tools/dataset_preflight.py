"""Read-only dataset inventory and explicitly approved, journaled in-place renames."""
from __future__ import annotations

import argparse
import csv
import ctypes
import errno
import hashlib
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from oscardp.shots.discovery import IMDB_PATTERN, VIDEO_EXTENSIONS

DEFAULT_MOVIES = Path('/media/sdu/SiyaoDu/datasets/oscar_movie')
DEFAULT_SCRIPTS = Path('/media/sdu/SiyaoDu/datasets/oscar_script')
DEFAULT_REPORTS = Path('reports/dataset_inventory')
DEFAULT_TITLES = Path('data/metadata/movie_titles.json')
SCRIPT_EXTENSIONS = {'.pdf'}  # Confirmed from current screenplay directory.
PLAN_FIELDS = ['movie_id', 'file_type', 'original_path', 'proposed_path', 'existing_filename',
               'proposed_filename', 'matching_evidence', 'collision_status', 'classification',
               'approval_status', 'review_notes', 'source_dev', 'source_ino', 'source_size', 'source_mtime_ns']
AUDIT_FIELDS = ['movie_id', 'movie_title', 'movie_year', 'video_status', 'video_path', 'srt_status',
                'srt_path', 'screenplay_status', 'screenplay_path', 'video_candidates', 'srt_candidates',
                'screenplay_candidates', 'video_match', 'srt_match', 'screenplay_match',
                'manual_video', 'manual_script', 'manual_discrepancy', 'manual_review_required']


@dataclass(frozen=True)
class Movie:
    movie_id: str
    title: str | None
    year: int | None
    manual_video: object
    manual_script: object


def _id(value: object) -> str:
    match = IMDB_PATTERN.fullmatch(str(value or '').strip())
    if not match:
        raise ValueError(f'Invalid FilmId: {value!r}')
    return match.group(1).lower()


def read_workbook(path: Path) -> dict[str, Movie]:
    if not path.is_file():
        raise FileNotFoundError(f'Workbook not found: {path}')
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book:
            iterator = sheet.iter_rows(values_only=True)
            for header in iterator:
                fields = {str(value).strip().casefold(): index for index, value in enumerate(header) if value is not None}
                if not {'filmid', 'film', 'year'} <= fields.keys():
                    continue
                movies: dict[str, Movie] = {}
                def cell(row, name):
                    index = fields.get(name)
                    return row[index] if index is not None and index < len(row) else None
                for row in iterator:
                    raw_id = cell(row, 'filmid')
                    if raw_id is None or str(raw_id).strip() == '':
                        continue
                    movie_id = _id(raw_id)
                    if movie_id in movies:
                        raise ValueError(f'Duplicate FilmId in workbook: {movie_id}')
                    raw_title = cell(row, 'film')
                    title = str(raw_title).strip() if raw_title is not None and str(raw_title).strip() else None
                    raw_year = cell(row, 'year')
                    year = int(raw_year) if raw_year is not None and str(raw_year).strip() else None
                    movies[movie_id] = Movie(movie_id, title, year, cell(row, 'video'), cell(row, 'script'))
                if not movies:
                    raise ValueError('Workbook has no movie records')
                return movies
        raise ValueError('No worksheet contains FilmId, Film and Year columns')
    finally:
        book.close()


def _normal(value: str) -> str:
    value = unicodedata.normalize('NFKD', value).casefold()
    return ''.join(char for char in value if char.isalnum() and not unicodedata.combining(char))


def _year(value: str) -> int | None:
    match = re.search(r'(?<!\d)(19\d{2}|20\d{2})(?!\d)', value)
    return int(match.group(1)) if match else None


def _title_key(stem: str) -> str:
    year = re.search(r'(?<!\d)(?:19|20)\d{2}(?!\d)', stem)
    if year:
        leading = _normal(stem[:year.start()])
        return leading or _normal(stem[year.end():])
    return _normal(stem)


def scan(root: Path, extensions: set[str]) -> list[Path]:
    if not root.is_dir():
        raise FileNotFoundError(f'Dataset directory not found: {root}')
    result = []
    for directory, dirs, files in os.walk(root, followlinks=False):
        dirs[:] = [name for name in dirs if not name.startswith('.') and not (Path(directory) / name).is_symlink()]
        for name in files:
            path = Path(directory) / name
            if name.startswith('.') or path.is_symlink() or path.suffix.lower() not in extensions:
                continue
            result.append(path.resolve())
    return sorted(result)


def file_id(path: Path, root: Path) -> tuple[str | None, str]:
    match = IMDB_PATTERN.search(path.name)
    if match:
        return match.group(1).lower(), 'filename_id'
    for parent in path.relative_to(root).parents[:-1]:
        match = IMDB_PATTERN.search(parent.name)
        if match:
            return match.group(1).lower(), 'parent_id'
    return None, ''


def _candidate(path: Path, root: Path, movies: dict[str, Movie]) -> tuple[str | None, str]:
    movie_id, method = file_id(path, root)
    if movie_id:
        return movie_id, method
    title = _title_key(path.stem)
    year = _year(path.stem)
    exact = [movie for movie in movies.values() if movie.title and _normal(movie.title) == title]
    if year is not None:
        exact = [movie for movie in exact if movie.year == year]
        if len(exact) == 1:
            return exact[0].movie_id, 'title_year'
    if len(exact) == 1 and year is None:
        return exact[0].movie_id, 'title_only_review'
    return None, 'unmatched'


def _status(candidates: list[tuple[Path, str]]) -> tuple[str, str, str]:
    if not candidates:
        return 'MISSING', '', ''
    if len(candidates) > 1:
        return 'MULTIPLE', '', ','.join(sorted({method for _, method in candidates}))
    path, method = candidates[0]
    if method in {'title_only_review', 'parent_id_without_video'}:
        return 'AMBIGUOUS', '', method
    return 'FOUND', str(path), method


def _manual_flag(value: object) -> bool | None:
    if value is None or str(value).strip() == '':
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    text = str(value).strip().casefold()
    if text in {'yes', 'y', 'true', '1', 'found', 'exists', '有', '是'}:
        return True
    if text in {'no', 'n', 'false', '0', 'missing', 'none', '无', '否'}:
        return False
    return None


def inventory(workbook: Path, movies_root: Path, scripts_root: Path):
    movies = read_workbook(workbook)
    videos = scan(movies_root, VIDEO_EXTENSIONS)
    srts = scan(movies_root, {'.srt'})
    scripts = scan(scripts_root, SCRIPT_EXTENSIONS)
    video_map: dict[str, list[tuple[Path, str]]] = defaultdict(list)
    script_map: dict[str, list[tuple[Path, str]]] = defaultdict(list)
    extras = []
    for kind, paths, root, target in [('video', videos, movies_root, video_map),
                                      ('screenplay', scripts, scripts_root, script_map)]:
        for path in paths:
            movie_id, method = _candidate(path, root, movies)
            if movie_id in movies:
                target[movie_id].append((path, method))
            else:
                if not file_id(path, root)[0]:
                    title = _title_key(path.stem)
                    plausible = [movie for movie in movies.values() if movie.title and _normal(movie.title) == title]
                    for movie in plausible:
                        target[movie.movie_id].append((path, 'title_only_review'))
                extras.append({'file_type': kind, 'path': str(path), 'reason': f'{method}: {movie_id or "no match"}'})
    rows = []
    srt_map: dict[str, list[tuple[Path, str]]] = {}
    used_srts = set()
    for movie_id, movie in movies.items():
        vc = video_map[movie_id]
        sc = script_map[movie_id]
        vs, vp, vm = _status(vc)
        ps, pp, pm = _status(sc)
        related_dirs = {path.parent for path, _ in vc}
        subtitles = [(path, 'same_video_directory') for path in srts if path.parent in related_dirs]
        if not vc:
            subtitles = [(path, 'parent_id_without_video') for path in srts
                         if file_id(path, movies_root)[0] == movie_id]
        srt_map[movie_id] = subtitles
        used_srts.update(path for path, _ in subtitles)
        ss, sp, sm = _status(subtitles)
        discrepancies = []
        for label, manual, actual in [('video', movie.manual_video, vs), ('script', movie.manual_script, ps)]:
            flag = _manual_flag(manual)
            if flag is not None and flag != (actual == 'FOUND'):
                discrepancies.append(label)
        rows.append({'movie_id': movie_id, 'movie_title': movie.title or '', 'movie_year': movie.year or '',
                     'video_status': vs, 'video_path': vp, 'srt_status': ss, 'srt_path': sp,
                     'screenplay_status': ps, 'screenplay_path': pp, 'video_candidates': len(vc),
                     'srt_candidates': len(subtitles), 'screenplay_candidates': len(sc),
                     'video_match': vm, 'srt_match': sm, 'screenplay_match': pm,
                     'manual_video': '' if movie.manual_video is None else str(movie.manual_video),
                     'manual_script': '' if movie.manual_script is None else str(movie.manual_script),
                     'manual_discrepancy': ','.join(discrepancies),
                     'manual_review_required': vs != 'FOUND' or ss != 'FOUND' or ps != 'FOUND'
                     or vm not in {'filename_id', 'parent_id'} or bool(discrepancies)})
    for path in srts:
        if path not in used_srts:
            movie_id, method = file_id(path, movies_root)
            extras.append({'file_type': 'srt', 'path': str(path), 'reason': f'{method or "unmatched"}: {movie_id or "no match"}'})
    for movie_id, matches in srt_map.items():
        if len(matches) > 1 or any(method == 'parent_id_without_video' for _, method in matches):
            extras.extend({'file_type': 'srt', 'path': str(path), 'reason': f'{movie_id}: ambiguous SRT ownership'}
                          for path, _ in matches)
    for kind, mapping in [('video', video_map), ('screenplay', script_map)]:
        for movie_id, matches in mapping.items():
            if matches and (len(matches) > 1 or matches[0][1] == 'title_only_review'):
                extras.extend({'file_type': kind, 'path': str(path), 'reason': f'{movie_id}: ambiguous candidates'} for path, _ in matches)
    unexpected = 0
    for item in extras:
        root = scripts_root if item['file_type'] == 'screenplay' else movies_root
        owner, _ = file_id(Path(item['path']), root)
        if owner not in movies and 'ambiguous' not in item['reason']:
            unexpected += 1
    summary = {'total_reference_movies': len(movies), 'unexpected_files': unexpected,
               'unmatched_or_ambiguous_files': len(extras),
               'complete_sets': sum(row['video_status'] == row['srt_status'] == row['screenplay_status'] == 'FOUND' for row in rows),
               'manual_discrepancies': sum(bool(row['manual_discrepancy']) for row in rows)}
    for field, prefix in [('video_status', 'videos'), ('srt_status', 'srts'), ('screenplay_status', 'screenplays')]:
        counts = Counter(row[field] for row in rows)
        for status in ('FOUND', 'MISSING', 'AMBIGUOUS', 'MULTIPLE'):
            summary[f'{prefix}_{status.lower()}'] = counts[status]
    return movies, rows, extras, summary, video_map, srt_map


def _write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + '\n', encoding='utf-8')


def write_audit(reports: Path, rows: list[dict], extras: list[dict], summary: dict) -> None:
    _write_csv(reports / 'inventory_audit.csv', AUDIT_FIELDS, rows)
    _write_csv(reports / 'unmatched_files.csv', ['file_type', 'path', 'reason'], extras)
    _write_json(reports / 'inventory_summary.json', summary)
    print(json.dumps(summary, sort_keys=True))
    print(json.dumps({'inventory_audit': str(reports / 'inventory_audit.csv'),
                      'unmatched_files': str(reports / 'unmatched_files.csv'),
                      'inventory_summary': str(reports / 'inventory_summary.json')}, sort_keys=True))


def plan_rename(movies: dict[str, Movie], rows: list[dict], video_map, srt_map) -> list[dict]:
    by_id = {row['movie_id']: row for row in rows}
    plan = []
    for movie_id in movies:
        row = by_id[movie_id]
        videos = video_map[movie_id]
        srts = srt_map[movie_id]
        eligible = (row['video_status'] == row['srt_status'] == 'FOUND' and len(videos) == len(srts) == 1
                    and videos[0][1] in {'filename_id', 'parent_id'})
        for kind, candidates in [('video', videos), ('srt', srts)]:
            for source, evidence in candidates:
                destination = source.with_name(movie_id + (source.suffix if kind == 'video' else '.srt'))
                same = source == destination
                names = {path.name.casefold() for path in source.parent.iterdir() if path != source}
                conflict = destination.exists() and not same or destination.name.casefold() in names
                classification = ('ALREADY_STANDARDISED' if same else 'CONFLICT' if conflict else
                                  'SAFE' if eligible and len(candidates) == 1 else 'REVIEW_REQUIRED')
                stat = source.stat()
                plan.append({'movie_id': movie_id, 'file_type': kind, 'original_path': str(source),
                             'proposed_path': str(destination), 'existing_filename': source.name,
                             'proposed_filename': destination.name, 'matching_evidence': evidence,
                             'collision_status': 'COLLISION' if conflict else 'CLEAR',
                             'classification': classification, 'approval_status': 'PENDING',
                             'review_notes': '', 'source_dev': stat.st_dev, 'source_ino': stat.st_ino,
                             'source_size': stat.st_size, 'source_mtime_ns': stat.st_mtime_ns})
    destinations = Counter(item['proposed_path'].casefold() for item in plan if item['classification'] == 'SAFE')
    for item in plan:
        if item['classification'] == 'SAFE' and destinations[item['proposed_path'].casefold()] > 1:
            item['classification'] = 'CONFLICT'
            item['collision_status'] = 'PLANNED_COLLISION'
    return plan


def _rename_noreplace(source: Path, destination: Path) -> None:
    libc = ctypes.CDLL(None, use_errno=True)
    result = libc.renameat2(-100, os.fsencode(source), -100, os.fsencode(destination), 1)
    if result != 0:
        code = ctypes.get_errno()
        raise OSError(code, os.strerror(code), str(destination))


def _journal(path: Path, entry: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('a', encoding='utf-8') as handle:
        handle.write(json.dumps(entry, sort_keys=True) + '\n')
        handle.flush()
        os.fsync(handle.fileno())


def _events(path: Path) -> dict[str, dict]:
    states = {}
    if path.exists():
        for line in path.read_text(encoding='utf-8').splitlines():
            event = json.loads(line)
            states[event['operation_id']] = event
    return states


def _identity(path: Path, row: dict) -> bool:
    if not path.is_file():
        return False
    stat = path.stat()
    return (stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns) == tuple(int(row[key]) for key in ('source_dev', 'source_ino', 'source_size', 'source_mtime_ns'))


def _stamp(path: Path) -> dict[str, int]:
    stat = path.stat()
    return {'dev': stat.st_dev, 'ino': stat.st_ino, 'size': stat.st_size, 'mtime_ns': stat.st_mtime_ns}


def _matches_stamp(path: Path, stamp: dict) -> bool:
    return path.is_file() and _stamp(path) == {key: int(value) for key, value in stamp.items()}


def _matches_prepared_destination(path: Path, row: dict) -> bool:
    if not path.is_file():
        return False
    stat = path.stat()
    return (stat.st_dev, stat.st_size, stat.st_mtime_ns) == tuple(
        int(row[key]) for key in ('source_dev', 'source_size', 'source_mtime_ns')
    )


def _approved_rows(path: Path) -> list[dict]:
    with path.open(newline='', encoding='utf-8') as handle:
        rows = list(csv.DictReader(handle))
    if not rows or not all(set(PLAN_FIELDS) <= row.keys() for row in rows):
        raise ValueError('Invalid or empty approved plan')
    return [row for row in rows if row['approval_status'] == 'APPROVED']


def _operation_id(row: dict) -> str:
    return hashlib.sha256((row['movie_id'] + '\0' + row['original_path'] + '\0' + row['proposed_path']).encode()).hexdigest()


def _validate_operation(row: dict) -> tuple[Path, Path]:
    if row['classification'] != 'SAFE' or row['collision_status'] != 'CLEAR':
        raise ValueError('Only SAFE, collision-free operations may be approved')
    source, destination = Path(row['original_path']), Path(row['proposed_path'])
    if not source.is_absolute() or not destination.is_absolute() or source.parent != destination.parent or source == destination:
        raise ValueError('Rename paths must be distinct absolute paths in the same directory')
    if row['file_type'] not in {'video', 'srt'} or row['matching_evidence'] not in {'filename_id', 'parent_id', 'same_video_directory'}:
        raise ValueError('Unverified file ownership')
    if row['file_type'] == 'video' and source.suffix.lower() not in VIDEO_EXTENSIONS:
        raise ValueError('Unsupported video extension')
    expected = row['movie_id'] + (source.suffix if row['file_type'] == 'video' else '.srt')
    if destination.name != expected or not IMDB_PATTERN.fullmatch(row['movie_id']):
        raise ValueError('Destination does not follow the naming convention')
    if row['file_type'] == 'srt' and source.suffix.lower() != '.srt':
        raise ValueError('Source is not an SRT')
    owner = IMDB_PATTERN.search(source.name) or IMDB_PATTERN.search(source.parent.name)
    if owner is None or owner.group(1).lower() != row['movie_id']:
        raise ValueError('Source IMDb ID cannot be verified')
    siblings = [path for path in source.parent.iterdir() if path.is_file() and not path.is_symlink()]
    if sum(path.suffix.lower() in VIDEO_EXTENSIONS for path in siblings) != 1 or sum(path.suffix.lower() == '.srt' for path in siblings) != 1:
        raise ValueError('Movie directory no longer has exactly one video and one SRT')
    if source.is_symlink() or destination.is_symlink():
        raise ValueError('Symlink rename is forbidden')
    if any(name.casefold() == destination.name.casefold() for name in os.listdir(source.parent) if name != source.name):
        raise FileExistsError(f'Destination collision: {destination}')
    return source, destination


def approve_safe_plan(plan: list[dict], expected_ids: set[str]) -> list[dict]:
    by_movie: dict[str, set[str]] = defaultdict(set)
    sources: set[str] = set()
    destinations: set[str] = set()
    approved = []
    for row in plan:
        movie_id, kind, classification = row['movie_id'], row['file_type'], row['classification']
        if movie_id not in expected_ids or kind in by_movie[movie_id]:
            raise ValueError(f'Unexpected or duplicate planned movie file: {movie_id} {kind}')
        by_movie[movie_id].add(kind)
        if classification not in {'SAFE', 'ALREADY_STANDARDISED'}:
            raise ValueError(f'Unresolved rename operation: {movie_id} {kind} {classification}')
        if classification == 'SAFE':
            source, destination = _validate_operation(row)
            if not _identity(source, row):
                raise RuntimeError(f'Source changed since planning: {source}')
            if str(source) in sources or str(destination).casefold() in destinations:
                raise ValueError('Conflicting source or destination in rename plan')
            sources.add(str(source))
            destinations.add(str(destination).casefold())
            approved.append({**row, 'approval_status': 'APPROVED'})
    if set(by_movie) != expected_ids or any(kinds != {'video', 'srt'} for kinds in by_movie.values()):
        raise ValueError('Rename plan does not cover exactly one video and SRT per reference movie')
    return approved


def apply_rename(plan_path: Path, journal: Path, confirm: bool) -> dict:
    if not confirm:
        raise ValueError('apply-rename requires --confirm')
    rows = _approved_rows(plan_path)
    if not rows:
        raise ValueError('No APPROVED operations in plan')
    ids = [_operation_id(row) for row in rows]
    if len(ids) != len(set(ids)):
        raise ValueError('Duplicate operations in plan')
    destinations = [row['proposed_path'].casefold() for row in rows]
    if len(destinations) != len(set(destinations)):
        raise ValueError('Approved plan has duplicate destinations')
    states = _events(journal)
    for row, operation_id in zip(rows, ids):
        state = states.get(operation_id, {}).get('state')
        source, destination = Path(row['original_path']), Path(row['proposed_path'])
        if state in {'DONE', 'ROLLED_BACK'}:
            continue
        if state == 'PREPARED' and _matches_prepared_destination(destination, row) and not source.exists():
            continue
        _validate_operation(row)
        if not _identity(source, row):
            raise RuntimeError(f'Source changed since planning: {source}')
    done = 0
    for row, operation_id in zip(rows, ids):
        source, destination = Path(row['original_path']), Path(row['proposed_path'])
        previous = states.get(operation_id, {})
        if previous.get('state') == 'DONE':
            if not _matches_stamp(destination, previous['after']) or source.exists():
                raise RuntimeError(f'Journal and filesystem disagree: {operation_id}')
            continue
        if previous.get('state') == 'ROLLED_BACK':
            raise RuntimeError('Rolled-back operation cannot be applied again from the same journal')
        if previous.get('state') == 'PREPARED' and _matches_prepared_destination(destination, row) and not source.exists():
            _journal(journal, {'operation_id': operation_id, 'state': 'DONE', 'row': row, 'after': _stamp(destination)})
            done += 1
            continue
        _validate_operation(row)
        if not _identity(source, row):
            raise RuntimeError(f'Source changed since planning: {source}')
        _journal(journal, {'operation_id': operation_id, 'state': 'PREPARED', 'row': row})
        _rename_noreplace(source, destination)
        _journal(journal, {'operation_id': operation_id, 'state': 'DONE', 'row': row, 'after': _stamp(destination)})
        done += 1
    return {'approved': len(rows), 'applied_or_recovered': done, 'journal': str(journal)}


def verify_journal(journal: Path, rollback: bool = False, confirm: bool = False) -> dict:
    if not journal.is_file():
        raise FileNotFoundError(f'Journal not found: {journal}')
    if rollback and not confirm:
        raise ValueError('Rollback requires --confirm')
    states = _events(journal)
    checked = rolled_back = 0
    for operation_id, event in reversed(list(states.items())):
        row = event['row']
        source, destination = Path(row['original_path']), Path(row['proposed_path'])
        state = event['state']
        if state == 'PREPARED':
            if _matches_prepared_destination(destination, row) and not source.exists():
                state = 'DONE'
                event = {'operation_id': operation_id, 'state': 'DONE', 'row': row, 'after': _stamp(destination)}
                _journal(journal, event)
            elif _identity(source, row) and not destination.exists():
                state = 'NOT_APPLIED'
            else:
                raise RuntimeError(f'Interrupted operation has conflicting paths: {operation_id}')
        if rollback and state == 'DONE':
            if not _matches_stamp(destination, event['after']) or source.exists():
                raise RuntimeError(f'Rollback conflict: {operation_id}')
            if any(name.casefold() == source.name.casefold() for name in os.listdir(source.parent) if name != destination.name):
                raise FileExistsError(f'Rollback destination collision: {source}')
            _rename_noreplace(destination, source)
            _journal(journal, {'operation_id': operation_id, 'state': 'ROLLED_BACK', 'row': row, 'after': _stamp(source)})
            rolled_back += 1
        elif not rollback and state == 'DONE' and (not _matches_stamp(destination, event['after']) or source.exists()):
            raise RuntimeError(f'Renamed file changed or disappeared: {destination}')
        elif not rollback and state == 'ROLLED_BACK' and (not _matches_stamp(source, event['after']) or destination.exists()):
            raise RuntimeError(f'Rolled-back file changed or disappeared: {source}')
        checked += 1
    return {'checked': checked, 'rolled_back': rolled_back, 'journal': str(journal)}


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description=__doc__)
    commands = root.add_subparsers(dest='command', required=True)
    for name in ('audit', 'plan-rename'):
        command = commands.add_parser(name)
        command.add_argument('--workbook', type=Path, required=True)
        command.add_argument('--movies-root', type=Path, default=DEFAULT_MOVIES)
        command.add_argument('--scripts-root', type=Path, default=DEFAULT_SCRIPTS)
        command.add_argument('--reports-dir', type=Path, default=DEFAULT_REPORTS)
        if name == 'plan-rename':
            command.add_argument('--titles-output', type=Path, default=DEFAULT_TITLES)
    apply = commands.add_parser('apply-rename')
    apply.add_argument('--approved-plan', type=Path, required=True)
    apply.add_argument('--journal', type=Path, default=DEFAULT_REPORTS / 'rename_journal.jsonl')
    apply.add_argument('--confirm', action='store_true')
    verify = commands.add_parser('verify')
    verify.add_argument('--journal', type=Path)
    verify.add_argument('--workbook', type=Path)
    verify.add_argument('--movies-root', type=Path, default=DEFAULT_MOVIES)
    verify.add_argument('--scripts-root', type=Path, default=DEFAULT_SCRIPTS)
    verify.add_argument('--reports-dir', type=Path, default=DEFAULT_REPORTS)
    verify.add_argument('--rollback', action='store_true')
    verify.add_argument('--confirm', action='store_true')
    return root


def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    try:
        if args.command in {'audit', 'plan-rename'}:
            movies, rows, extras, summary, videos, srts = inventory(args.workbook, args.movies_root, args.scripts_root)
            write_audit(args.reports_dir, rows, extras, summary)
            if args.command == 'plan-rename':
                plan = plan_rename(movies, rows, videos, srts)
                _write_csv(args.reports_dir / 'rename_plan.csv', PLAN_FIELDS, plan)
                if any(not movie.title for movie in movies.values()):
                    raise ValueError('Cannot create title mapping: workbook contains missing titles')
                titles = {movie_id: movie.title for movie_id, movie in movies.items() if movie.title}
                _write_json(args.titles_output, titles)
                print(json.dumps({'rename_plan': dict(Counter(item['classification'] for item in plan)),
                                  'titles': len(titles), 'missing_titles': len(movies) - len(titles),
                                  'rename_plan_path': str(args.reports_dir / 'rename_plan.csv'),
                                  'titles_path': str(args.titles_output)}, sort_keys=True))
        elif args.command == 'apply-rename':
            print(json.dumps(apply_rename(args.approved_plan, args.journal, args.confirm)))
        else:
            if not args.journal and not args.workbook:
                raise ValueError('verify requires --journal or --workbook')
            if args.rollback and not args.journal:
                raise ValueError('Rollback requires --journal')
            if args.journal:
                print(json.dumps(verify_journal(args.journal, args.rollback, args.confirm)))
            if args.workbook:
                _, rows, extras, summary, _, _ = inventory(args.workbook, args.movies_root, args.scripts_root)
                write_audit(args.reports_dir, rows, extras, summary)
        return 0
    except (OSError, ValueError, RuntimeError) as exc:
        print(f'error: {exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
